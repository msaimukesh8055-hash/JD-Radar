"""Populate one catalog slide with data and save it as a standalone .pptx.

Usage:
    python3 populate.py <sample_index> <data.json> <output.pptx>

data.json: { "role_name": "text" }  or  { "role_name": ["line1","line2",...] }
For table roles:  { "role_name": [["r0c0","r0c1"], ["r1c0","r1c1"], ...] }
For chart roles:   { "role_name": {"categories": [...], "series": {"name": [..]}} }
"""
import json
import sys
import shutil

from pptx import Presentation

from pptx_utils import (
    find_shape_by_id,
    set_paragraph_lines,
    set_table_cell_text,
    replace_chart_data,
    remove_other_slides,
)

CATALOG_PATH = "/home/user/slide-smith/catalog/slide-index.json"
SAMPLE_DECK = "/home/user/slide-smith/templates/sample-20.pptx"


def load_catalog_entry(sample_index):
    with open(CATALOG_PATH) as f:
        catalog = json.load(f)
    for entry in catalog["slides"]:
        if entry["sample_index"] == sample_index:
            return entry
    raise ValueError(f"No catalog entry for sample_index={sample_index}")


def apply_data(slide, entry, data):
    fillable_by_role = {item["role"]: item for item in entry.get("fillable", [])}

    for role, value in data.items():
        item = fillable_by_role.get(role)
        if item is None:
            print(f"  [skip] unknown role '{role}' for slide '{entry['label']}'")
            continue

        shape = find_shape_by_id(slide.shapes, item["shape_id"])
        if shape is None:
            print(f"  [skip] shape_id {item['shape_id']} not found")
            continue

        if "table" in item:
            rows = value
            for r_i, row_vals in enumerate(rows):
                for c_i, cell_text in enumerate(row_vals):
                    cell = shape.table.cell(r_i, c_i)
                    lines = cell_text if isinstance(cell_text, list) else [cell_text]
                    set_table_cell_text(cell, lines)
            continue

        if "chart_type" in item:
            categories = value["categories"]
            series = [(name, vals) for name, vals in value["series"].items()]
            replace_chart_data(shape, categories, series)
            continue

        lines = value if isinstance(value, list) else [value]
        set_paragraph_lines(shape.text_frame, lines)


def main():
    sample_index = int(sys.argv[1])
    data_path = sys.argv[2]
    out_path = sys.argv[3]

    entry = load_catalog_entry(sample_index)
    with open(data_path) as f:
        data = json.load(f)

    shutil.copy(SAMPLE_DECK, out_path)
    prs = Presentation(out_path)
    slide = remove_other_slides(prs, sample_index - 1)

    print(f"Populating '{entry['label']}' (category: {entry['category']})")
    apply_data(slide, entry, data)

    prs.save(out_path)
    print(f"Saved {out_path}")


if __name__ == "__main__":
    main()
