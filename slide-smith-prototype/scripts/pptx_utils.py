"""Shared helpers for finding shapes and writing text/table/chart data
into a slide while preserving the template's existing formatting."""
import copy
from pptx.shapes.group import GroupShape
from pptx.util import Pt
from pptx.chart.data import CategoryChartData


def find_shape_by_id(shapes, target_id):
    """Recursively search shapes (including nested groups) for a shape_id."""
    for shape in shapes:
        if shape.shape_id == target_id:
            return shape
        if isinstance(shape, GroupShape):
            found = find_shape_by_id(shape.shapes, target_id)
            if found is not None:
                return found
    return None


def set_paragraph_lines(text_frame, lines):
    """Write `lines` (list of str) into text_frame, one per paragraph,
    reusing existing paragraphs/runs (and their formatting) where possible.
    Extra paragraphs beyond len(lines) are removed; if more lines than
    paragraphs exist, the last paragraph is cloned to make room."""
    paragraphs = list(text_frame.paragraphs)

    # grow: clone the last paragraph's XML for any extra lines needed
    while len(paragraphs) < len(lines):
        last_p = paragraphs[-1]._p
        new_p = copy.deepcopy(last_p)
        last_p.addnext(new_p)
        from pptx.text.text import _Paragraph
        paragraphs.append(_Paragraph(new_p, text_frame))

    # shrink: drop trailing paragraphs we don't need
    for p in paragraphs[len(lines):]:
        p._p.getparent().remove(p._p)
    paragraphs = paragraphs[:len(lines)]

    for p, line in zip(paragraphs, lines):
        runs = p.runs
        if not runs:
            run = p.add_run()
            run.text = line
            continue
        runs[0].text = line
        for extra in runs[1:]:
            extra._r.getparent().remove(extra._r)


def set_table_cell_text(cell, lines):
    set_paragraph_lines(cell.text_frame, lines)


def replace_chart_data(chart_shape, categories, series):
    """series: list of (name, values) tuples."""
    chart_data = CategoryChartData()
    chart_data.categories = categories
    for name, values in series:
        chart_data.add_series(name, values)
    chart_shape.chart.replace_data(chart_data)


def remove_other_slides(prs, keep_index_0based):
    """Strip every slide except the one at keep_index_0based, returning
    the (now sole) slide."""
    from pptx.oxml.ns import qn
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    for i in range(len(slides) - 1, -1, -1):
        if i == keep_index_0based:
            continue
        rId = slides[i].get(qn('r:id'))
        prs.part.drop_rel(rId)
        xml_slides.remove(slides[i])
    return prs.slides[0]
