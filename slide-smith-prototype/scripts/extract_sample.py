"""Extract a subset of slides from the master deck into a small sample deck,
preserving layouts, masters, images and chart parts (by deleting the rest)."""
import shutil
import sys
from pptx import Presentation
from pptx.oxml.ns import qn

SRC = "/root/.claude/uploads/6e012faa-d745-5442-9771-094dc642f2a3/6dcea620-16x9_Timesaver_PPT.pptx"
DST = "/home/user/slide-smith/templates/sample-20.pptx"

# 1-based slide indices in the master deck, in ascending order
KEEP = [1, 4, 10, 12, 19, 20, 21, 24, 26, 57, 62, 74, 86, 98, 123, 138, 159, 214, 221, 256]


def remove_slide(prs, slide_index_0based):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    rId = slides[slide_index_0based].get(qn('r:id'))
    prs.part.drop_rel(rId)
    xml_slides.remove(slides[slide_index_0based])


def main():
    shutil.copy(SRC, DST)
    prs = Presentation(DST)

    total = len(prs.slides)
    keep_set = set(KEEP)
    remove_indices = [i for i in range(total) if (i + 1) not in keep_set]

    for i in sorted(remove_indices, reverse=True):
        remove_slide(prs, i)

    # strip the now-dangling p14:sectionLst (lives in p:presentation/p:extLst)
    root = prs.part._element
    ext_lst = root.find(qn('p:extLst'))
    if ext_lst is not None:
        root.remove(ext_lst)

    prs.save(DST)
    print(f"Saved {DST} with {len(Presentation(DST).slides)} slides "
          f"(expected {len(KEEP)})")


if __name__ == "__main__":
    main()
